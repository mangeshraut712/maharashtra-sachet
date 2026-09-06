#!/usr/bin/env python3
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "web" / "pitch.pptx"
MEDIA = ROOT / "web" / "pitch-media"
GREEN = RGBColor(0x25, 0x45, 0x63)
ACCENT = RGBColor(0xA4, 0x4B, 0x13)
BODY = RGBColor(0x26, 0x2B, 0x31)
MUTED = RGBColor(0x58, 0x61, 0x6B)


def add_text(shape, text, size=18, bold=False, color=BODY, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def bullets(shape, lines, size=18):
    tf = shape.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = BODY
        run.font.name = "Calibri"


def title_slide(prs, kicker, title, body):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    k = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(12), Inches(0.4))
    add_text(k, kicker, 14, True, ACCENT)
    t = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12), Inches(1.6))
    add_text(t, title, 36, True, GREEN)
    b = slide.shapes.add_textbox(Inches(0.6), Inches(2.8), Inches(12), Inches(3.4))
    bullets(b, body, 20)
    return slide


def picture_slide(prs, kicker, title, image, caption):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    k = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.35))
    add_text(k, kicker, 13, True, ACCENT)
    t = slide.shapes.add_textbox(Inches(0.5), Inches(0.55), Inches(12), Inches(0.6))
    add_text(t, title, 26, True, GREEN)
    slide.shapes.add_picture(str(image), Inches(0.5), Inches(1.25), width=Inches(12.3))
    c = slide.shapes.add_textbox(Inches(0.5), Inches(6.95), Inches(12.3), Inches(0.4))
    add_text(c, caption, 13, False, MUTED)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)

    title_slide(
        prs,
        "CODEX BUILD HOUSE · PUNE",
        "Maharashtra Civic Alerts",
        [
            "Independent English/Marathi hub for official Maharashtra public warnings.",
            "Not a government website. Not emergency dispatch. Call 112 in danger.",
            "Empty bulletin is not an all-clear.",
            "Builder: Mangesh Raut  ·  Live: Cloudflare Workers + D1",
        ],
    )
    title_slide(
        prs,
        "PROBLEM",
        "Residents should not guess which feed is live",
        [
            "Official sources are scattered across agencies.",
            "An empty map looks safe when a source is stale or disabled.",
            "Place names are ambiguous (Navi Mumbai spans two districts).",
            "Warnings must stay attributed, not rewritten by a model.",
        ],
    )
    title_slide(
        prs,
        "SOLUTION",
        "A place-aware bulletin that states its coverage",
        [
            "CAP lifecycle: actual, update, cancel, expiry.",
            "District search and aliases: Sawantwadi, Marunji, Navi Mumbai.",
            "Source-health strip before anyone reads “no alerts”.",
            "12 citizen-service directories labelled live vs directory-only.",
            "Labelled local demo fixtures that never enter production.",
        ],
    )
    picture_slide(
        prs,
        "LIVE PRODUCT",
        "Production hub on Cloudflare Workers",
        MEDIA / "desktop-home.png",
        "https://maharashtra-sachet.mangeshraut712.workers.dev",
    )
    picture_slide(
        prs,
        "SEARCH HONESTY",
        "Places map to districts, not invented village polygons",
        MEDIA / "desktop-search-navi-mumbai.png",
        "Navi Mumbai returns separate Raigad and Thane candidates.",
    )
    picture_slide(
        prs,
        "SOURCE HEALTH",
        "Read freshness before sharing an empty bulletin",
        MEDIA / "desktop-source-health.png",
        "CWC and CPCB stay disabled until a reviewed adapter exists.",
    )
    picture_slide(
        prs,
        "MARATHI",
        "Same bulletin, resident language",
        MEDIA / "desktop-marathi.png",
        "Original warning text is kept. Language fallback is labelled.",
    )
    title_slide(
        prs,
        "LIMITS",
        "What this project does not do",
        [
            "No government affiliation, cell broadcast or 112 dispatch.",
            "No complete village or ward map.",
            "No background Web Push after the tab closes.",
            "No generative rewrite of emergency instructions.",
            "The repository predates the event. Event work: honesty, realtime ingest, place search, labelled demo.",
        ],
    )
    title_slide(
        prs,
        "LINKS",
        "Review these surfaces",
        [
            "Live demo: https://maharashtra-sachet.mangeshraut712.workers.dev",
            "Pitch: https://maharashtra-sachet.mangeshraut712.workers.dev/pitch.html",
            "GitHub: https://github.com/mangeshraut712/maharashtra-sachet",
            "Local labelled demo: npm run demo  →  http://127.0.0.1:8799",
            "Phone: +91 7276819090  ·  Email: mbr63drexel@gmail.com",
        ],
    )
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
